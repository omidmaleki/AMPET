import os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\Lucifer\Documents\Default Project\AMPET_OpenCode\Technical Architecture"
MD = os.path.join(BASE, "Technical-Architecture.md")
DOCX_OUT = os.path.join(BASE, "Technical-Architecture.docx")

PRIMARY = RGBColor(0x0F, 0x5D, 0x4A)
GOLD = RGBColor(0xD4, 0xAF, 0x37)
BLACK = RGBColor(0x22, 0x22, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PRIMARY_HEX = "0F5D4A"
ALT_HEX = "EEF3F1"

def set_shading(cell, color_hex):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    shd.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shd)

def add_para(doc, text, size=11, bold=False, color=BLACK, align="RIGHT", space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align=="RIGHT" else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Vazirmatn'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), 'Vazirmatn')
    rFonts.set(qn('w:hAnsi'), 'Vazirmatn')
    rFonts.set(qn('w:cs'), 'Vazirmatn')
    return p

def add_heading(doc, text, level):
    if level == 1:
        return add_para(doc, text, size=18, bold=True, color=PRIMARY)
    if level == 2:
        return add_para(doc, text, size=14, bold=True, color=PRIMARY)
    return add_para(doc, text, size=12, bold=True, color=GOLD)

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = 'Vazirmatn'; run.font.size = Pt(9); run.font.bold = True
        run.font.color.rgb = WHITE
        set_shading(cell, PRIMARY_HEX)
    for ri, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[ri+1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = 'Vazirmatn'; run.font.size = Pt(8)
            run.font.color.rgb = BLACK
            if ri % 2 == 0:
                set_shading(cell, ALT_HEX)
    doc.add_paragraph()
    return table

def parse_md(path):
    with open(path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    return lines

def is_table_sep(line):
    return bool(re.match(r'^\s*\|?[\s:\-|]+\|?\s*$', line)) and '-' in line

def build_docx():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Vazirmatn'
    style.font.size = Pt(11)
    sectPr = doc.sections[0]._sectPr
    rtl = OxmlElement('w:bidi')
    rtl.set(qn('w:val'), 'true')
    sectPr.append(rtl)

    # Cover
    for _ in range(5):
        doc.add_paragraph()
    add_para(doc, "AMPET", size=40, bold=True, color=PRIMARY, align="CENTER")
    add_para(doc, "معماری فنی — Technical Architecture", size=20, bold=True, color=GOLD, align="CENTER")
    add_para(doc, "نسخه ۱٫۰ — ۱۴۰۵/۰۷/۱۴", size=12, color=BLACK, align="CENTER")
    add_para(doc, "رنگ برند: سبز 0F5D4A • طلایی D4AF37", size=11, color=BLACK, align="CENTER")
    add_para(doc, "مرجع: project-document.md + وایرفریم‌های ۲۸گانه", size=10, color=BLACK, align="CENTER")
    doc.add_page_break()

    lines = parse_md(MD)
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # headings
        m = re.match(r'^(#{1,3})\s+(.*)', line)
        if m:
            level = len(m.group(1))
            add_heading(doc, m.group(2).strip(), level)
            i += 1
            continue
        # table
        if line.strip().startswith('|') and i+1 < len(lines) and is_table_sep(lines[i+1]):
            headers = [c.strip() for c in line.strip().strip('|').split('|')]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            add_table(doc, headers, rows)
            continue
        # blockquote
        if line.strip().startswith('>'):
            add_para(doc, line.strip().lstrip('>').strip(), size=10, bold=False, color=PRIMARY)
            i += 1
            continue
        # bullets / numbered
        if re.match(r'^\s*([-\*]|\d+[.)])\s+', line):
            t = re.sub(r'^\s*([-\*]|\d+[.)])\s+', '• ', line).strip()
            add_para(doc, t, size=10)
            i += 1
            continue
        # hr
        if re.match(r'^\s*---+\s*$', line):
            add_para(doc, "— — —", size=9, color=GOLD, align="CENTER")
            i += 1
            continue
        # code fence skip markers
        if line.strip().startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                add_para(doc, lines[i], size=9, color=BLACK)
                i += 1
            i += 1
            continue
        add_para(doc, line.strip(), size=10)
        i += 1

    doc.save(DOCX_OUT)
    print(f"DOCX saved: {DOCX_OUT}")

if __name__ == "__main__":
    build_docx()
